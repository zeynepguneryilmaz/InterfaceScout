from __future__ import annotations

import json
import shutil
import urllib.request
from pathlib import Path

from openpyxl import Workbook

from interface_engine import analyze_all_maps

ROOT = Path(__file__).resolve().parents[1]
EXAMPLES = ROOT / "examples"

# Only structures and conditions described in the manuscript are generated.
CASES = [
    # Parameter-robustness development panel (pH 7.4, 150 mM, 298 K).
    {"pdb": "1CRN", "chain": "A", "pH": 7.4, "map": "anionic", "folder": "1CRN"},
    {"pdb": "1R69", "chain": "A", "pH": 7.4, "map": "anionic", "folder": "1R69"},
    {"pdb": "1SHG", "chain": "A", "pH": 7.4, "map": "anionic", "folder": "1SHG"},
    {"pdb": "2PPN", "chain": "A", "pH": 7.4, "map": "anionic", "folder": "2PPN"},
    {"pdb": "4AKE", "chain": "A", "pH": 7.4, "map": "anionic", "folder": "4AKE"},
    {"pdb": "1OMP", "chain": "A", "pH": 7.4, "map": "anionic", "folder": "1OMP"},
    {"pdb": "1TIM", "chain": "A,B", "pH": 7.4, "map": "anionic", "folder": "1TIM"},
    {"pdb": "5CSC", "chain": "B", "pH": 7.4, "map": "anionic", "folder": "5CSC"},
    # Literature-comparison cases.
    {"pdb": "4F5S", "chain": "A", "pH": 8.0, "map": "anionic", "folder": "4F5S"},
    {"pdb": "2VUF", "chain": "A", "pH": 6.2, "map": "anionic", "folder": "2VUF"},
    {"pdb": "2CBA", "chain": "A", "pH": 6.3, "map": "anionic", "folder": "2CBA/pH_6.3"},
    {"pdb": "2CBA", "chain": "A", "pH": 8.3, "map": "anionic", "folder": "2CBA/pH_8.3"},
    {"pdb": "1UBQ", "chain": "A", "pH": 7.7, "map": "anionic", "folder": "1UBQ"},
    {"pdb": "1JNJ", "chain": "A", "pH": 7.0, "map": "anionic", "folder": "1JNJ"},
    {"pdb": "3ECA", "chain": "A,B,C,D", "pH": 7.5, "map": "cationic", "folder": "3ECA"},
    {"pdb": "1SUV", "chain": "C,E", "pH": 7.4, "map": "hydrophobic", "folder": "1SUV"},
    {"pdb": "1TGU", "chain": "A,B,C,D", "pH": 7.4, "map": "hydrophobic", "folder": "1TGU"},
]


def fetch_pdb(pdb_id: str) -> str:
    with urllib.request.urlopen(f"https://files.rcsb.org/download/{pdb_id}.pdb", timeout=60) as response:
        return response.read().decode("utf-8", errors="replace")


def flatten(obj, prefix=""):
    rows = []
    if obj is None:
        return [(prefix, "")]
    if isinstance(obj, list):
        return [(prefix, "; ".join(json.dumps(v, separators=(",", ":")) if isinstance(v, (dict, list)) else str(v) for v in obj))]
    if isinstance(obj, dict):
        for key, value in obj.items():
            rows.extend(flatten(value, f"{prefix}.{key}" if prefix else key))
        return rows
    return [(prefix, obj)]


def residue_rows(data):
    all_rows = {}
    for key in data["map_order"]:
        for residue in data["maps"][key].get("residues", []):
            row = all_rows.setdefault(residue["key"], {
                "ResidueKey": residue["key"],
                "Chain": residue.get("chain"),
                "ResidueNumber": residue.get("res_seq"),
                "InsertionCode": residue.get("icode", ""),
                "ResidueName": residue.get("res_name"),
                "scRSA": residue.get("scrsa"),
            })
            row[f"{key}_local_score"] = residue.get("local_score")
            row[f"{key}_propensity_0_100"] = residue.get("propensity")
            row[f"{key}_multiscale_persistence"] = residue.get("multiscale_persistence")
    return sorted(all_rows.values(), key=lambda r: (str(r.get("Chain") or ""), int(r.get("ResidueNumber") or 0), str(r.get("InsertionCode") or "")))


def patch_rows(data):
    rows = []
    for key in data["map_order"]:
        map_data = data["maps"][key]
        for patch in map_data.get("patches", []):
            rows.append({
                "MapKey": key,
                "MapLabel": map_data["label"],
                "DisplayRank": patch.get("display_rank"),
                "ParetoFront": patch.get("pareto_front"),
                "Classification": patch.get("classification", ""),
                "CenterKey": patch.get("center_key"),
                "MemberCount": patch.get("n_members"),
                "ChemistrySeedCount": patch.get("n_chemistry_seeds"),
                "Chains": ";".join(patch.get("chains", [])),
                "ResidueSpanMin": patch.get("residue_span_min"),
                "ResidueSpanMax": patch.get("residue_span_max"),
                "Diameter_A": patch.get("diameter_A"),
                "ChemistrySupport": patch.get("chemistry_support"),
                "MeanAccessibility": patch.get("mean_accessibility"),
                "PatchCoherence": patch.get("patch_coherence"),
                "OrientationCoherence": patch.get("orientation_coherence"),
                "Members": ";".join(patch.get("members", [])),
                "ChemistrySeeds": ";".join(patch.get("seed_members", [])),
            })
    return rows


def patch_member_rows(data):
    rows = []
    for key in data["map_order"]:
        map_data = data["maps"][key]
        for patch in map_data.get("patches", []):
            seeds = set(patch.get("seed_members", []))
            for residue in patch.get("member_residues", []):
                rows.append({
                    "MapKey": key,
                    "MapLabel": map_data["label"],
                    "PatchRank": patch.get("display_rank"),
                    "ParetoFront": patch.get("pareto_front"),
                    "ResidueKey": residue.get("key", ""),
                    "Chain": residue.get("chain"),
                    "ResidueNumber": residue.get("res_seq"),
                    "InsertionCode": residue.get("icode", ""),
                    "ResidueName": residue.get("res_name"),
                    "IsChemistrySeed": residue.get("key") in seeds,
                })
    return rows


def add_sheet(workbook, name, rows):
    worksheet = workbook.create_sheet(name)
    if not rows:
        worksheet.append(["Info"])
        worksheet.append(["No data"])
        return
    headers = list(rows[0].keys())
    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(header) for header in headers])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for index, header in enumerate(headers, 1):
        max_len = len(str(header))
        for cells in worksheet.iter_cols(min_col=index, max_col=index, min_row=2, max_row=min(worksheet.max_row, 201)):
            for cell in cells:
                max_len = max(max_len, len(str(cell.value or "")))
        worksheet.column_dimensions[worksheet.cell(1, index).column_letter].width = min(28, max_len + 2)


def write_excel(data, path):
    workbook = Workbook()
    workbook.remove(workbook.active)
    run = [
        {"Parameter": "engine", "Value": data.get("engine")},
        {"Parameter": "model_version", "Value": data.get("version")},
        {"Parameter": "public_version", "Value": data.get("public_version", "1.0-publication")},
    ]
    for section in ["input", "structure_preparation", "method", "scope", "diagnostics", "composite_policy"]:
        for key, value in flatten(data.get(section), section):
            run.append({"Parameter": key, "Value": value})
    add_sheet(workbook, "Run_Method", run)
    definitions = []
    for key in data["map_order"]:
        map_data = data["maps"][key]
        definitions.append({
            "MapKey": key,
            "MapLabel": map_data["label"],
            "Description": map_data["description"],
            "InternalChemistryKey": map_data["internal_chemistry_key"],
            "PatchCount": map_data["n_patches"],
            "ParetoFront1Count": map_data["n_pareto_primary_patches"],
        })
    add_sheet(workbook, "Map_Definitions", definitions)
    add_sheet(workbook, "Residue_Values", residue_rows(data))
    add_sheet(workbook, "Patches", patch_rows(data))
    add_sheet(workbook, "Patch_Members", patch_member_rows(data))
    workbook.save(path)


def rewrite_bfactor(raw_pdb, rows, label):
    lookup = {}
    for residue in rows:
        key = (str(residue.get("chain") or "").strip(), int(residue.get("res_seq")), str(residue.get("icode") or "").strip())
        lookup[key] = max(0.0, min(100.0, float(residue.get("propensity") or 0.0)))
    output = []
    for line in raw_pdb.replace("\r", "").split("\n"):
        if line.startswith(("ATOM  ", "HETATM")):
            padded = line.ljust(66)
            chain = padded[21:22].strip()
            try:
                resi = int(padded[22:26].strip())
            except ValueError:
                resi = 0
            icode = padded[26:27].strip()
            bfactor = f"{lookup.get((chain, resi, icode), 0.0):6.2f}"
            line = padded[:60] + bfactor + padded[66:]
        output.append(line)
    remarks = [
        "REMARK 900 INTERFACESCOUT B-FACTOR TRACK",
        f"REMARK 900 MAP: {label}",
        "REMARK 900 FIELD: canonical propensity; NORMALIZED RANGE 0-100",
        "REMARK 900 B-FACTOR VALUES ARE INTERFACESCOUT SCORES, NOT EXPERIMENTAL CRYSTALLOGRAPHIC B-FACTORS",
    ]
    return "\n".join(remarks + output)


def main():
    if EXAMPLES.exists():
        shutil.rmtree(EXAMPLES)
    for case in CASES:
        print("Generating", case)
        raw = fetch_pdb(case["pdb"])
        data = analyze_all_maps(
            pdb_text=raw,
            pdb_id=case["pdb"],
            chain=case["chain"],
            pH=case["pH"],
            ionic_mM=150.0,
            temp_K=298.0,
        )
        data["public_version"] = "1.0-publication"
        outdir = EXAMPLES / case["folder"]
        outdir.mkdir(parents=True, exist_ok=True)
        suffix = f"_pH{case['pH']}" if case["pdb"] == "2CBA" else ""
        write_excel(data, outdir / f"InterfaceScout_{case['pdb']}{suffix}_all_results.xlsx")
        map_data = data["maps"][case["map"]]
        pdb_output = rewrite_bfactor(raw, map_data.get("residues", []), map_data["label"])
        (outdir / f"InterfaceScout_{case['pdb']}{suffix}_{case['map']}_Bfactor.pdb").write_text(pdb_output, encoding="utf-8")


if __name__ == "__main__":
    main()
